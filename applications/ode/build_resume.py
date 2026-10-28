from pathlib import Path
from html import escape
import json
import re
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, KeepTogether
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

for alias,filename in [('Helvetica','LiberationSans-Regular.ttf'),('Helvetica-Bold','LiberationSans-Bold.ttf'),('Helvetica-Oblique','LiberationSans-Italic.ttf')]:
    pdfmetrics.registerFont(TTFont(alias,'/usr/share/fonts/truetype/liberation/'+filename))

OUT = Path(__file__).resolve().parent
STEM = 'Jacob_Nikolau_Ode_Resume'
LINKS = {
    'github.com/NYCandrun': 'https://github.com/NYCandrun',
    'Customer research tools': 'https://github.com/jacobn-a11y/Transcript-Project-',
    'Company research workflows': 'https://github.com/NYCandrun/genesis/tree/template/clean-research-workflow',
    'Additional projects': 'https://github.com/NYCandrun/resume/blob/main/applications/ode/Project_Portfolio.md',
    'AssistRx': 'https://www.mosaicapp.com/customers/assistrx',
    'CetraRuddy': 'https://www.mosaicapp.com/customers/cetraruddy',
    'Stahly Engineering': 'https://www.mosaicapp.com/customers/stahly-engineering',
    'CESO': 'https://www.mosaicapp.com/customers/ceso',
    'Clark Patterson Lee': 'https://www.mosaicapp.com/customers/cpl',
    'Vocon': 'https://www.mosaicapp.com/customers/vocon',
    'The State of AI in A&E 2026': 'https://www.mosaicapp.com/whitepapers/state-of-ai-in-a-e-2026',
    'HR Dive': 'https://www.hrdive.com/news/learning-talent-development-pros-foresee-demise-in-person-corporate-training/635000/',
    'TD Magazine': 'https://www.td.org/content/td-magazine/say-goodbye-to-in-person-training',
    'Harvard Business Review': 'https://hbr.org/2021/01/how-companies-are-using-vr-to-develop-employees-soft-skills',
    'Forbes': 'https://www.forbes.com/sites/jeannemeister/2021/01/14/10-hr-trends-for-the-next-normal-of-work/',
}

LINK_PATTERN = re.compile('|'.join(re.escape(label) for label in sorted(LINKS, key=len, reverse=True)))

def linked_parts(txt):
    """Split once so multiple links retain their order without rewriting URLs."""
    cursor = 0
    for match in LINK_PATTERN.finditer(txt):
        if match.start() > cursor:
            yield txt[cursor:match.start()], None
        yield match.group(), LINKS[match.group()]
        cursor = match.end()
    if cursor < len(txt):
        yield txt[cursor:], None

blocks = [
('name', 'JACOB NIKOLAU'),
('tagline', 'Product Marketing | AI & Technical B2B Software'),
('contact', 'New York, NY | jacobnikolau767@gmail.com | 310-985-4638 | github.com/NYCandrun'),
('section', 'SUMMARY'),
('body', 'Product marketer with eight years in enterprise and AI software. I research industries and work directly with customers to understand which business problems a product can solve. I help executives choose where to focus and help sales teams make the case to enterprise buyers. I’ve been the first marketer at two startups. At NovoEd, I designed ROI research using six executive interviews and a 127-person survey. At Mosaic, I built Claude workflows that cut case-study production from 50 hours to five. I also build AI prototypes and test models to understand their capabilities and limits.'),
('section', 'PROFESSIONAL EXPERIENCE'),
('role', 'Mosaic | Director of Marketing'),
('meta', 'AI operations and resource-management software | May 2025 – Jul 2026'),
('bullet', 'Built Claude skills to turn sales-call objections, competitor mentions, and customer outcomes into reusable evidence, reducing case-study production from 50 hours to 5 using a RevOps-built Gong/Salesforce integration.'),
('bullet', 'Directed AI-assisted research across 13 architecture, engineering, and construction sub-industries; recommended sector-specific positioning and built a sales-development playbook grounded in observable prospect signals.'),
('bullet', 'Managed approximately $130K in quarterly media and sponsorship spend; shifted budget toward webinars after finding that their cost per sales-accepted opportunity was 83% lower than paid digital.'),
('bullet', 'Founded the 21-member Manage Better Board for architecture and engineering operators; ran six executive roundtables and industry sessions where customers shared workflows and results.'),
('bullet', 'Positioned Chello, Mosaic’s AI resource-management agent, around operations leaders’ priorities; created sales talk tracks distinguishing live capabilities from roadmap commitments.'),
('bullet', 'Worked with AssistRx, CetraRuddy, Stahly Engineering, CESO, Clark Patterson Lee, and Vocon on approved case studies with ROI calculations, giving customers and sales teams evidence to explain business value.'),
('bullet', 'Rebuilt attribution using an LLM-driven workflow that matched attendee emails across 2,190 sales calls to deal contacts and sourcing teams.'),
('bullet', 'Led a three-person marketing team plus agencies as Mosaic’s first marketer; grew marketing-sourced opportunities from 68 to 238 per quarter across research, content, events, and campaigns.'),
('bullet', 'Designed and fielded a 138-person survey and co-authored The State of AI in A&E 2026, examining AI adoption among architecture and engineering professionals.'),
('role', 'BitcoinComputer.io | Head of Growth, Product Marketing (Consultant)'),
('meta', 'Bitcoin smart-contract platform | Aug 2024 – May 2025'),
('bullet', 'Served as the first and sole marketer for a pre-revenue platform; wrote a messaging framework defining target audiences, value proposition, differentiators, and competitive positioning.'),
('bullet', 'Translated the platform’s architecture into website content for technical and business evaluation; researched competing protocols and regulatory considerations for skill gaming to clarify product tradeoffs.'),
('role', 'NovoEd | Head of Product Marketing'),
('meta', 'Enterprise learning platform | Apr 2023 – Aug 2024'),
('bullet', 'Worked with Sales, Customer Experience, and Professional Services to turn customer use cases into deal-specific sales materials. Supported wins and expansions with Delta, 3M, Centene, Marriott, Capital Group, and IMD Business School.'),
('bullet', 'Owned positioning and sales enablement for the AI Translation launch, partnering with Sales and Customer Success on rollout. The feature reached about 100,000 learners in its first year and drove over $1M in expansion revenue across CEMEX and Marriott.'),
('bullet', 'Built messaging and features-to-value frameworks for business buyers and technical evaluators in 12–24-month sales cycles, including regulated accounts such as Goldman Sachs and Capital Group.'),
('bullet', 'Built a customer-story and reference program across 15 enterprise customers and managed customer legal approvals; arranged for Marriott’s CHRO to present the engagement at Josh Bersin’s conference.'),
('bullet', 'Ran analyst relations with Gartner and Forrester; earned a Craig Weiss Top 10 Learning Systems placement.'),
('pagebreak', ''),
('section', 'PROFESSIONAL EXPERIENCE, CONTINUED'),
('role', 'NovoEd | Marketing Strategy Manager'),
('meta', 'Oct 2021 – Apr 2023'),
('bullet', 'Designed ROI research with Cascade Insights, including six executive interviews and a 127-person survey; translated customer-reported outcomes into a shared value framework for Sales and Customer Success.'),
('bullet', 'Developed industry-specific positioning and packaging across manufacturing, hospitality, financial services, technology, healthcare, and energy. Co-built a market-sizing model estimating a ~$9B TAM and ~$1.4B mid-market opportunity.'),
('bullet', 'Helped customers explain their experiences through webinars and ebooks that were the first conversion point for 33% of marketing-qualified leads and 21% of sales-qualified leads.'),
('bullet', 'Authored research with Executive Networks based on a survey of 515 learning leaders across five countries, earning coverage in HR Dive and TD Magazine.'),
('role', 'Mursion | Manager, Growth Marketing'),
('meta', 'AI simulation training | Jul 2020 – Sep 2021'),
('bullet', 'Developed go-to-market strategies across technology, healthcare, hospitality, consumer financial services, and call-center use cases. Built the healthcare vertical from scratch and landed Duke University Health System; ranked first on the team in new opportunities created and was fastest to reach $3M in attributed pipeline.'),
('bullet', 'Helped enterprise customers explain their use of AI simulation through webinars and live demonstrations; developed a co-branded webinar program with Google.'),
('bullet', 'Authored research with Future Workplace based on a survey of 300 corporate learning leaders; Harvard Business Review built an article around the findings, and Forbes cited the study.'),
('role', 'Feedonomics | Growth Engineer'),
('meta', 'Commerce data and feed-integration software | Aug 2018 – Jul 2020'),
('bullet', 'Owned first-touch qualification for enterprise and mid-market prospects; used sales-conversation patterns to improve messaging and worked with RevOps on lead scoring, funnel tracking, and response time.'),
('section', 'SELECTED APPLIED AI WORK'),
('role', 'Independent Projects | 2026'),
('body', 'Built self-directed tools and prototypes with Claude Code and Codex to understand how products work, make their value tangible, and turn customer conversations into usable evidence.'),
('bullet', 'Company research workflows: Source collection, citation checks, and independent review for company briefs, with controls that stop completion when required evidence is missing.'),
('bullet', 'Model evaluation: Vision-language-model benchmark comparing small models; found that prompt formatting could shift accuracy by about 30 percentage points on a hallucination benchmark.'),
('bullet', 'Compute marketplace: Fixed-price AI-job workflows and a 517-company competitive dossier testing the business case for verified results and buyer recourse.'),
('bullet', 'Customer research tools: Desktop Gong/Grain transcript downloader and Markdown merger, plus case-study agents that organize conversations by account, preserve source attribution, and produce drafts and quote exports.'),
('body', 'Additional projects'),
('section', 'SKILLS & TOOLS'),
('body', 'Positioning, value frameworks, customer research and advocacy, market sizing, competitive analysis, AI workflows, model evaluation. Claude, Claude Code, Codex, Gong, Salesforce, HubSpot.'),
('section', 'EDUCATION'),
('body', 'University of California, Berkeley — BA, Geography; dual concentration in Earth & Planetary Science and Environmental Economics.'),
]

OUT.mkdir(parents=True, exist_ok=True)
md=[]
for typ, txt in blocks:
    if typ == 'pagebreak': continue
    txt = ''.join(f'[{part}]({url})' if url else part for part, url in linked_parts(txt))
    if typ == 'name': md.append('# '+txt)
    elif typ == 'section': md.append('## '+txt.title().replace('Ai', 'AI'))
    elif typ == 'role': md.append('### '+txt)
    elif typ == 'bullet': md.append('- '+txt)
    elif typ == 'meta': md.append('*'+txt+'*')
    else: md.append(txt)
(OUT/(STEM+'.md')).write_text('\n\n'.join(md)+'\n')

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.55); sec.bottom_margin=Inches(.55)
sec.left_margin=Inches(.65); sec.right_margin=Inches(.65)
normal=doc.styles['Normal']
normal.font.name='Calibri'; normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(3)
normal.paragraph_format.line_spacing=1.02
for typ,txt in blocks:
    if typ=='pagebreak': doc.add_page_break(); continue
    p=doc.add_paragraph()
    r=p.add_run(txt)
    if typ=='name':
        r.bold=True; r.font.size=Pt(22); p.paragraph_format.space_after=Pt(3)
    elif typ=='tagline':
        r.bold=True; r.font.size=Pt(10.5)
    elif typ=='contact':
        r.font.size=Pt(9.5); p.paragraph_format.space_after=Pt(5)
    elif typ=='section':
        r.bold=True; r.font.size=Pt(10)
        r.font.color.rgb=RGBColor.from_string('243B53')
        p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(3)
    elif typ=='role':
        r.bold=True; p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(2)
    elif typ=='meta':
        r.italic=True; r.font.size=Pt(9.5); p.paragraph_format.space_after=Pt(4)
    elif typ=='bullet':
        p.style=doc.styles['List Bullet']
        p.paragraph_format.left_indent=Inches(.13)
        p.paragraph_format.first_line_indent=Inches(-.13)
        p.paragraph_format.space_after=Pt(1.5)
    if typ in ('name','tagline','contact','section','role','meta'):
        p.paragraph_format.keep_with_next=True
    p.paragraph_format.widow_control=True
    if LINK_PATTERN.search(txt):
        formatting = deepcopy(r._r.rPr) if r._r.rPr is not None else None
        p._p.remove(r._r)
        for part,url in linked_parts(txt):
            if not url:
                run = p.add_run(part)
                if formatting is not None:
                    run._r.insert(0, deepcopy(formatting))
                continue
            link=OxmlElement('w:hyperlink')
            link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
            lr=OxmlElement('w:r')
            props=OxmlElement('w:rPr')
            size=OxmlElement('w:sz'); size.set(qn('w:val'),'19' if typ=='contact' else '21'); props.append(size)
            color=OxmlElement('w:color'); color.set(qn('w:val'),'243B53'); props.append(color)
            underline=OxmlElement('w:u'); underline.set(qn('w:val'),'single'); props.append(underline)
            lr.append(props)
            lt=OxmlElement('w:t'); lt.text=part; lr.append(lt)
            link.append(lr); p._p.append(link)
doc.core_properties.title='Jacob Nikolau — Ode, Associate Principal, AI Transformation'
doc.core_properties.author='Jacob Nikolau'
doc.save(OUT/(STEM+'.docx'))

styles={}
for typ in ('body','bullet','name','tagline','contact','section','role','meta'):
    styles[typ]=ParagraphStyle(typ,fontName='Helvetica',fontSize=10,leading=12,spaceAfter=3)
styles['name']=ParagraphStyle('name',fontName='Helvetica-Bold',fontSize=22,leading=25,spaceAfter=4)
styles['tagline']=ParagraphStyle('tagline',fontName='Helvetica-Bold',fontSize=10.2,leading=13,spaceAfter=4)
styles['contact']=ParagraphStyle('contact',fontName='Helvetica',fontSize=9.2,leading=11,spaceAfter=5)
styles['section']=ParagraphStyle('section',fontName='Helvetica-Bold',fontSize=9.5,leading=12,spaceBefore=6,spaceAfter=3,textColor=colors.HexColor('#243B53'),keepWithNext=True)
styles['role']=ParagraphStyle('role',fontName='Helvetica-Bold',fontSize=10,leading=12,spaceBefore=5,spaceAfter=2,keepWithNext=True)
styles['meta']=ParagraphStyle('meta',fontName='Helvetica-Oblique',fontSize=9.2,leading=11,spaceAfter=4,keepWithNext=True)
styles['bullet'].leftIndent=9
styles['bullet'].bulletIndent=0
styles['bullet'].spaceAfter=1.5
story=[]
for typ,txt in blocks:
    if typ=='pagebreak': story.append(PageBreak()); continue
    markup = ''.join(f'<a href="{escape(url, quote=True)}" color="#243B53"><u>{escape(part)}</u></a>' if url else escape(part) for part, url in linked_parts(txt))
    story.append(Paragraph(markup,styles[typ],bulletText='•' if typ=='bullet' else None))
def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont('Helvetica',8)
    canvas.setFillColor(colors.HexColor('#66788A'))
    canvas.drawString(46.8,23,'Jacob Nikolau')
    canvas.drawRightString(565.2,23,str(doc.page))
    canvas.restoreState()
pdf=SimpleDocTemplate(str(OUT/(STEM+'.pdf')),pagesize=letter,rightMargin=46.8,leftMargin=46.8,topMargin=36,bottomMargin=36,title='Jacob Nikolau — Ode, Associate Principal, AI Transformation',author='Jacob Nikolau')
pdf.build(story,onFirstPage=footer,onLaterPages=footer)
print('Created Markdown, DOCX and PDF resume drafts.')
print('Words:',sum(len(t.split()) for k,t in blocks))
